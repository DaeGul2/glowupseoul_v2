# -*- coding: utf-8 -*-
"""
Glow Up Seoul — v3 관리자 사용 가이드 (PPTX).
비전공자용. 실제 DB 값 예시 + 각 값이 고객 화면 어디에 보이는지.
DB 값은 2026-06-14 스냅샷. 다시 만들려면: python scripts/export_v3_admin_pptx.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------------- brand palette ----------------
INK      = RGBColor(0x18, 0x18, 0x1A)   # charcoal
PAPER    = RGBColor(0xFA, 0xFA, 0xF7)   # warm white
GOLD     = RGBColor(0xB8, 0x91, 0x6A)   # champagne
GOLD_DK  = RGBColor(0x8A, 0x6A, 0x4A)
MUTED    = RGBColor(0x6B, 0x6B, 0x6B)
LINE     = RGBColor(0xE4, 0xDF, 0xD6)
CARD     = RGBColor(0xF2, 0xEE, 0xE6)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
PINK     = RGBColor(0xC2, 0x4E, 0x7A)

FONT = "맑은 고딕"   # Malgun Gothic — Korean + Latin on Windows
TODAY = "2026-06-14"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def _set_run(r, size, color, bold=False, italic=False, font=FONT):
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color

def slide(bg=PAPER):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(1, 0, 0, SW, SH)  # rectangle bg
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s

def textbox(s, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if "space_before" in ln: p.space_before = Pt(ln["space_before"])
        p.space_after = Pt(ln.get("space_after", 2))
        if "line_spacing" in ln: p.line_spacing = ln["line_spacing"]
        runs = ln["runs"] if "runs" in ln else [ln]
        for rr in runs:
            run = p.add_run(); run.text = rr["t"]
            _set_run(run, rr.get("s", 16), rr.get("c", INK), rr.get("b", False), rr.get("i", False))
    return tb

def rect(s, x, y, w, h, fill=None, line=None, line_w=1.0):
    sh = s.shapes.add_shape(1, x, y, w, h)
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh

def eyebrow(s, txt, x=Inches(0.85), y=Inches(0.55)):
    textbox(s, x, y, Inches(11), Inches(0.4),
            [{"runs":[{"t":"✦  ","s":14,"c":GOLD,"b":True},{"t":txt.upper(),"s":13,"c":GOLD_DK,"b":True}]}])

def header(s, title, sub=None):
    eyebrow(s, "Glow Up Seoul · 관리자 가이드")
    textbox(s, Inches(0.85), Inches(0.95), Inches(11.6), Inches(1.0),
            [{"runs":[{"t":title,"s":34,"c":INK,"b":True}]}])
    rect(s, Inches(0.9), Inches(1.78), Inches(1.1), Pt(3), fill=GOLD)
    if sub:
        textbox(s, Inches(0.85), Inches(1.95), Inches(11.6), Inches(0.6),
                [{"runs":[{"t":sub,"s":14,"c":MUTED}]}])

def pagenum(s, n):
    textbox(s, Inches(11.8), Inches(7.0), Inches(1.3), Inches(0.35),
            [{"runs":[{"t":str(n),"s":11,"c":MUTED}]}], align=PP_ALIGN.RIGHT)

PAGE = [0]
def nextpage(s):
    PAGE[0]+=1; pagenum(s, PAGE[0])

# ---------------- DB snapshot ----------------
SURGERIES = [
    ("Rhinoplasty","Surgical reshaping of the nose.","₩4,000,000","Nose shape"),
    ("Double Eyelid Surgery","Creates a natural double-eyelid crease.","₩1,500,000","Double eyelid"),
    ("Ptosis Correction","Tightens the muscle that lifts a droopy eyelid.","₩2,000,000","Droopy eyelid"),
    ("Liposuction","Removes stubborn localized fat for a smoother line.","₩3,000,000","Localized fat"),
    ("Facelift","Surgically lifts deeper sagging for lasting rejuvenation.","₩12,000,000","Deep facial sagging"),
    ("Cleft Lip / Palate","Cleft lip and palate correction.","Consult","Cleft lip / palate"),
]
AREAS_NS = [("Skin","Pores, Rough texture, Pigmentation, Redness, Acne, Fine lines"),
            ("Eye area","Dark circles, Under-eye hollows"),
            ("Neck","Neck wrinkles"),
            ("Face & contour","Sagging & firmness, Volume loss, Face slimming")]
AREAS_S  = [("Eyes","Double eyelid, Droopy eyelid"),
            ("Nose","Nose shape"),
            ("Face lift","Deep facial sagging"),
            ("Body & liposuction","Localized fat"),
            ("Cleft lip / palate","Cleft lip / palate")]
TAGS = ("Lifting, Non-invasive, HIFU, Firming, RF, Skinbooster, Regeneration, Pores, Nose, Contour, "
        "Reconstructive, Special, Tightening, Injectable, Slimming, Wrinkles, Volume, Collagen, Laser, "
        "Pigmentation, Tone, Resurfacing, Scars, Facial, Hydration, Threads, IV, Brightening, Wellness, "
        "Under-eye, Eyes, Signature, Functional, Body")
REASON = ("Ulthera is the only HIFU device cleared to target the SMAS — the exact connective layer a "
          "surgeon lifts in a facelift. Heating that layer produces a genuine, non-surgical lift of the "
          "jaw and cheeks over 2–3 months.")

# ---------------- table helper ----------------
def add_table(s, x, y, w, headers, rows, col_w, header_fill=INK, fs=11, hfs=12, row_h=Inches(0.42)):
    nr, nc = len(rows)+1, len(headers)
    gf = s.shapes.add_table(nr, nc, x, y, w, row_h*nr)
    tbl = gf.table
    tbl.first_row = False; tbl.horz_banding = False
    total = sum(col_w)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Emu(int(int(w)*cw/total))
    # header
    for c, htext in enumerate(headers):
        cell = tbl.cell(0, c); cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
        cell.margin_left=Pt(7); cell.margin_right=Pt(5); cell.margin_top=Pt(3); cell.margin_bottom=Pt(3)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; run = p.add_run(); run.text = htext
        _set_run(run, hfs, WHITE, bold=True)
    # body
    for r, rowvals in enumerate(rows, start=1):
        for c, val in enumerate(rowvals):
            cell = tbl.cell(r, c)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if r % 2 else CARD
            cell.margin_left=Pt(7); cell.margin_right=Pt(5); cell.margin_top=Pt(3); cell.margin_bottom=Pt(3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            segs = val if isinstance(val, list) else [(val, False, INK)]
            for (t, b, col) in segs:
                run = p.add_run(); run.text = t; _set_run(run, fs, col, bold=b)
    return gf

def bullets(s, x, y, w, h, items, fs=15, gap=8):
    lines=[]
    for it in items:
        if isinstance(it, tuple):
            txt, lvl = it
        else:
            txt, lvl = it, 0
        if txt == "":
            lines.append({"runs":[{"t":" ","s":6}]}); continue
        if lvl == -1:  # section label
            lines.append({"space_before":10,"runs":[{"t":txt,"s":fs+1,"c":GOLD_DK,"b":True}]})
        elif lvl == 0:
            lines.append({"space_before":gap,"line_spacing":1.05,
                          "runs":[{"t":"•  ","s":fs,"c":GOLD,"b":True},{"t":txt,"s":fs,"c":INK}]})
        else:
            lines.append({"space_before":3,"line_spacing":1.05,
                          "runs":[{"t":"– ","s":fs-1,"c":MUTED},{"t":txt,"s":fs-1,"c":MUTED}]})
    textbox(s, x, y, w, h, lines)

# ======================================================================
# 1. TITLE
# ======================================================================
s = slide(INK)
rect(s, 0, Inches(3.05), SW, Pt(2), fill=GOLD)
textbox(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5),
        [{"runs":[{"t":"✦  GLOW UP SEOUL","s":16,"c":GOLD,"b":True}]}])
textbox(s, Inches(1.0), Inches(3.25), Inches(11.3), Inches(1.6),
        [{"runs":[{"t":"관리자 페이지 사용 가이드","s":46,"c":PAPER,"b":True}]}])
textbox(s, Inches(1.0), Inches(4.7), Inches(11.3), Inches(0.8),
        [{"runs":[{"t":"시술 · 수술 · 고민을 직접 추가하고 수정하는 법  —  비전공자용","s":18,"c":RGBColor(0xCB,0xC4,0xB6)}]}])
textbox(s, Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.5),
        [{"runs":[{"t":f"기준일 {TODAY}   ·   화면 글자는 모두 영어로 입력합니다","s":13,"c":MUTED}]}])

# ======================================================================
# 2. 목차
# ======================================================================
s = slide(); header(s, "목차", "아래 순서대로 보시면 됩니다")
toc = [
    "1.  시작하기 — 접속과 로그인",
    "2.  화면 구성 — 탭 3개",
    "3.  꼭 기억할 5가지",
    "4.  시술 등록 따라하기 (Ulthera 예시)",
    "5.  시술 필드 사전 — 칸마다 무슨 값인가",
    "6.  선택지 값 모음 (통증·회복·지속)",
    "7.  수술 관리",
    "8.  고민(Concerns) 관리",
    "9.  고민 연결 & 이유 — 매칭의 핵심",
    "10. 태그 · 사진 · 자주 묻는 질문",
]
# two columns
textbox(s, Inches(0.95), Inches(2.4), Inches(5.9), Inches(4.5),
        [{"space_before":12,"runs":[{"t":t,"s":18,"c":INK}]} for t in toc[:5]])
textbox(s, Inches(7.0), Inches(2.4), Inches(5.9), Inches(4.5),
        [{"space_before":12,"runs":[{"t":t,"s":18,"c":INK}]} for t in toc[5:]])
nextpage(s)

# ======================================================================
# 3. 시작하기
# ======================================================================
s = slide(); header(s, "1. 시작하기", "관리자 페이지란? 고객 사이트의 내용을 직접 고치는 곳입니다")
bullets(s, Inches(0.95), Inches(2.4), Inches(11.6), Inches(4.4), [
    ("들어가는 법", -1),
    ("주소창에 사이트 주소 뒤에 /admin 을 붙입니다.  예) glowupseoul.com/admin", 0),
    ("관리자 키(비밀번호)를 입력하고 Sign in 을 누릅니다.", 0),
    ("키는 담당 개발자에게 받으세요. 분실·유출 시 바로 알려 교체하세요.", 0),
    ("여기서 저장하면 고객 사이트에 곧바로 반영됩니다.", 0),
    ("", 0),
    ("이 가이드는", -1),
    ("관리자에서 직접 바꾸는 부분(시술/수술/고민)을 다룹니다.", 0),
    ("디자인에 박힌 고정 문구는 별도 파일(클라이언트 명세서 엑셀)에서 검수합니다.", 0),
])
nextpage(s)

# ======================================================================
# 4. 화면 구성 — 탭 3개
# ======================================================================
s = slide(); header(s, "2. 화면 구성 — 상단 탭 3개")
cards = [
    ("Treatments", "시술 (비수술)", "울쎄라 · 써마지 · 필러 · 레이저 등", GOLD),
    ("Surgeries", "수술", "코 · 눈 · 지방흡입 · 안면거상 등", GOLD_DK),
    ("고민 Concerns", "부위 · 고민", "챗봇 질문지와 추천의 바탕", INK),
]
cx = Inches(0.95); cw = Inches(3.75); gap = Inches(0.18)
for i,(t,sub,desc,accent) in enumerate(cards):
    x = Emu(int(cx) + i*(int(cw)+int(gap)))
    rect(s, x, Inches(2.6), cw, Inches(2.7), fill=WHITE, line=LINE, line_w=1)
    rect(s, x, Inches(2.6), cw, Pt(5), fill=accent)
    textbox(s, Emu(int(x)+Inches(0.3)), Inches(3.0), Emu(int(cw)-int(Inches(0.6))), Inches(2.2),[
        {"runs":[{"t":t,"s":22,"c":INK,"b":True}]},
        {"space_before":4,"runs":[{"t":sub,"s":15,"c":GOLD_DK,"b":True}]},
        {"space_before":10,"line_spacing":1.1,"runs":[{"t":desc,"s":14,"c":MUTED}]},
    ])
textbox(s, Inches(0.95), Inches(5.7), Inches(11.6), Inches(1.0),
        [{"runs":[{"t":"각 탭에서 New 버튼으로 새로 만들고, 목록의 카드를 눌러 수정합니다.","s":15,"c":INK}]}])
nextpage(s)

# ======================================================================
# 5. 꼭 기억할 5가지
# ======================================================================
s = slide(); header(s, "3. 꼭 기억할 5가지", "이것만 지키면 사고가 안 납니다")
five = [
    ("①", "모든 글자는 영어로", "고객이 외국인입니다. 한글 초안을 주면 영문으로 다듬어 드립니다."),
    ("②", "Visible 체크 = 공개", "‘Visible on the site’ 를 끄면 고객 사이트에서 사라집니다."),
    ("③", "Slug 는 바꾸지 않기", "영문 주소 키. 바꾸면 링크가 깨집니다. 보통 자동 생성이라 그대로 두세요."),
    ("④", "사진은 가능하면", "필수는 아니지만, 없으면 카드가 밋밋합니다."),
    ("⑤", "고민 연결이 매칭", "시술마다 어떤 고민에 좋은지 꼭 연결하세요. 추천이 여기서 나옵니다."),
]
y = Inches(2.45)
for num, t, d in five:
    rect(s, Inches(0.95), y, Inches(11.45), Inches(0.78), fill=WHITE, line=LINE, line_w=1)
    textbox(s, Inches(1.15), y, Inches(0.8), Inches(0.78),
            [{"runs":[{"t":num,"s":24,"c":GOLD,"b":True}]}], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, Inches(2.0), y, Inches(3.2), Inches(0.78),
            [{"runs":[{"t":t,"s":17,"c":INK,"b":True}]}], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, Inches(5.2), y, Inches(7.05), Inches(0.78),
            [{"runs":[{"t":d,"s":14,"c":MUTED}]}], anchor=MSO_ANCHOR.MIDDLE)
    y = Emu(int(y) + int(Inches(0.9)))
nextpage(s)

# ======================================================================
# 6. 시술 등록 따라하기
# ======================================================================
s = slide(); header(s, "4. 시술 등록 따라하기", "Treatments 탭 → ＋ New Treatment.  예시: Ulthera")
steps = [
    ("1","Name","시술 이름 (영문·필수)","Ulthera"),
    ("2","Short summary","한 문장 설명","Non-invasive HIFU lifting…"),
    ("3","Photo","대표 사진 업로드 (드래그/클릭)","시술 대표 이미지 1장"),
    ("4","Tags","키워드 — Enter 추가, × 삭제","Lifting / Non-invasive / HIFU"),
    ("5","Description","본문 전체 (오른쪽 미리보기)","## What to expect …"),
    ("6","Price (KRW)","참고 가격(숫자). 비우면 Consult","700000"),
    ("7","Price note","가격 옆 단서","from 300 shots (reference)"),
    ("8","How long it lasts","효과 지속 — 보기 선택","~1–2 years"),
    ("9","Pain / Recovery","통증·회복 — 보기 선택","Mild / Back to normal right away"),
    ("10","Good for / Note","추천 대상·유의점 한 줄씩","incision 없이 리프팅 …"),
    ("11","고민 연결 + 이유","매칭 핵심! 고민 고르고 이유 입력","Sagging & firmness → 이유"),
    ("12","Visible → Save","공개 체크 후 저장","✅ 체크 → Save"),
]
rows = [[ [(n,True,GOLD_DK)], [(f,True,INK)], [(w,False,MUTED)], [(ex,False,INK)] ] for n,f,w,ex in steps]
add_table(s, Inches(0.85), Inches(2.45), Inches(11.65),
          ["#","화면 항목","무엇을 하나요","예시 입력"], rows, [0.5,2.3,3.6,4.0], fs=11, hfs=12, row_h=Inches(0.39))
nextpage(s)

# ======================================================================
# 7. 시술 필드 사전 (the money slide)
# ======================================================================
s = slide(); header(s, "5. 시술 필드 사전", "각 칸이 무슨 값이고, 고객 화면 어디에 보이는지")
fd = [
    ("Name","시술 이름(영문)","Ulthera, Rhinoplasty","카드 제목·상세 제목·챗봇 추천","필수"),
    ("Short summary","한 문장 설명","Non-invasive HIFU lifting…","카드/상세 부제·챗봇 설명","권장"),
    ("Photo","대표 사진","(업로드 이미지)","카드 썸네일·상세 큰 사진","권장"),
    ("Tags","키워드 칩","Lifting, HIFU","카드 하단 칩·상세 상단","선택"),
    ("Description","본문(Markdown)","제목·목록·굵게","상세 본문","권장"),
    ("Price (KRW)","참고 가격","700000 → ₩700,000","카드 from…·상세 Price","선택"),
    ("Price note","가격 단서","from 300 shots","상세 Price 옆","선택"),
    ("How long it lasts","효과 지속","~1–2 years","상세 Results last","선택"),
    ("Pain level","통증","Mild","상세 Comfort","선택"),
    ("Recovery time","회복","Back to normal…","카드 ⏱·상세 Downtime","선택"),
    ("Good for","추천 대상","incision 없이 리프팅","상세 Good for","선택"),
    ("Things to note","유의점","약간의 붓기","상세 Things to note","선택"),
    ("고민 연결","고민+이유(매칭)","Sagging & firmness …",[("상세 Helps with·Why it helps·챗봇",True,PINK)],"핵심"),
    ("Sort order","정렬(숫자)","1, 2, 3","목록 순서","선택"),
    ("Slug","URL 키(자동)","ulthera","주소·사진폴더","건들지마"),
    ("Visible","공개 여부","체크=보임",[("끄면 사이트에서 사라짐",True,GOLD_DK)],"중요"),
]
rows = []
for name,what,ex,where,req in fd:
    where_seg = where if isinstance(where, list) else [(where,False,INK)]
    rows.append([[(name,True,INK)],[(what,False,MUTED)],[(ex,False,INK)], where_seg, [(req,True,GOLD_DK)]])
add_table(s, Inches(0.6), Inches(2.3), Inches(12.15),
          ["화면 항목","무슨 값","실제 예시","고객 화면 위치","필수?"],
          rows, [1.7,2.2,2.6,3.4,1.2], fs=10, hfs=11, row_h=Inches(0.275))
nextpage(s)

# ======================================================================
# 8. 선택지 값 모음
# ======================================================================
s = slide(); header(s, "6. 선택지 값 모음", "버튼으로 고르면 됩니다 — 참고용")
def minicol(x, title, pairs, accent):
    rect(s, x, Inches(2.5), Inches(3.75), Inches(3.9), fill=WHITE, line=LINE, line_w=1)
    rect(s, x, Inches(2.5), Inches(3.75), Pt(4), fill=accent)
    lines=[{"runs":[{"t":title,"s":17,"c":INK,"b":True}]}]
    for a,b in pairs:
        lines.append({"space_before":9,"runs":[{"t":a+"  ","s":13,"c":GOLD_DK,"b":True},{"t":"→ "+b,"s":13,"c":MUTED}]})
    textbox(s, Emu(int(x)+int(Inches(0.28))), Inches(2.8), Inches(3.2), Inches(3.5), lines)
minicol(Inches(0.95),"효과 지속",[("temporary","Temporary"),("months_3_6","3–6개월"),("months_6_12","6–12개월"),("year_1_2","1–2년"),("years_2_plus","2년 이상"),("semi_permanent","반영구"),("permanent","영구")],GOLD)
minicol(Inches(4.88),"통증 / 회복",[("soft","거의 없음 🙂"),("mild","약간 😐"),("hard","있음 😣"),("immediate","바로 일상 ⚡"),("1_2_days","1–2일 🌙"),("1_week_plus","1주 이상 🗓️")],GOLD_DK)
minicol(Inches(8.8),"고민 관련도",[("primary","주된 해결 (기본)"),("secondary","보조적으로 도움")],INK)
nextpage(s)

# ======================================================================
# 9. 수술 관리
# ======================================================================
s = slide(); header(s, "7. 수술 관리", "입력 방법은 시술과 100% 동일 · Surgeries 탭에서")
rows = [[[(n,True,INK)],[(d,False,MUTED)],[(p,False,INK)],[(c,False,GOLD_DK)]] for n,d,p,c in SURGERIES]
add_table(s, Inches(0.85), Inches(2.4), Inches(11.65),
          ["수술 이름","한 줄 설명","참고가격","연결된 고민"], rows, [2.6,4.7,1.7,2.6], fs=11, hfs=12, row_h=Inches(0.42))
textbox(s, Inches(0.95), Inches(6.45), Inches(11.6), Inches(0.8),
        [{"runs":[{"t":"⚠ 현재 수술은 고민 ‘연결’은 됐지만 ‘이유’가 비어 있습니다. 이유를 채우면 상세의 Why it helps·챗봇 설명이 나옵니다.","s":13,"c":PINK,"b":True}]}])
nextpage(s)

# ======================================================================
# 10. 고민 관리
# ======================================================================
s = slide(); header(s, "8. 고민(Concerns) 관리", "부위(Area) → 그 안에 세부 고민. 챗봇 질문·추천의 바탕")
# left: NS, right: S
def area_block(x, title, data, accent):
    rect(s, x, Inches(2.4), Inches(5.7), Inches(4.4), fill=WHITE, line=LINE, line_w=1)
    rect(s, x, Inches(2.4), Inches(5.7), Pt(4), fill=accent)
    lines=[{"runs":[{"t":title,"s":16,"c":INK,"b":True}]}]
    for a,c in data:
        lines.append({"space_before":9,"runs":[{"t":a,"s":14,"c":GOLD_DK,"b":True}]})
        lines.append({"space_before":1,"line_spacing":1.05,"runs":[{"t":c,"s":12,"c":MUTED}]})
    textbox(s, Emu(int(x)+int(Inches(0.3))), Inches(2.7), Inches(5.1), Inches(4.0), lines)
area_block(Inches(0.95),"시술(비수술)용 부위 → 고민", AREAS_NS, GOLD)
area_block(Inches(6.85),"수술용 부위 → 고민", AREAS_S, GOLD_DK)
nextpage(s)

# ======================================================================
# 11. 고민 연결 & 이유
# ======================================================================
s = slide(); header(s, "9. 고민 연결 & 이유 — 매칭의 핵심", "시술 편집 화면 아래 ‘이 시술이 해결하는 고민’")
bullets(s, Inches(0.95), Inches(2.35), Inches(11.6), Inches(2.0), [
    ("이 시술이 좋은 고민을 골라 연결합니다 (여러 개 가능).", 0),
    ("각 고민마다 ‘왜 이 시술이 좋은지’ 이유를 영어로 적습니다.", 0),
    ("이유는 → 상세 ‘Why it helps’ + 챗봇 추천의 ‘Why for OOO’ 에 그대로 보입니다.", 0),
    ("연결만 하고 이유를 비우면 추천엔 잡히지만 설명은 안 나옵니다.", 0),
])
rect(s, Inches(0.95), Inches(4.65), Inches(11.45), Inches(1.95), fill=CARD, line=GOLD, line_w=1)
textbox(s, Inches(1.25), Inches(4.85), Inches(10.9), Inches(1.7), [
    {"runs":[{"t":"실제 예시  ","s":14,"c":GOLD_DK,"b":True},{"t":"Ulthera → Sagging & firmness","s":14,"c":INK,"b":True}]},
    {"space_before":7,"line_spacing":1.12,"runs":[{"t":REASON,"s":13,"c":INK,"i":True}]},
])
nextpage(s)

# ======================================================================
# 12. 태그 · 사진 · FAQ
# ======================================================================
s = slide(); header(s, "10. 태그 · 사진 · 자주 묻는 질문")
# tags
rect(s, Inches(0.95), Inches(2.35), Inches(11.45), Inches(1.15), fill=WHITE, line=LINE, line_w=1)
textbox(s, Inches(1.2), Inches(2.5), Inches(11.0), Inches(0.95), [
    {"runs":[{"t":"태그(Tags)  ","s":15,"c":INK,"b":True},{"t":"카드 하단 칩(최대 3개) · 상세 상단 라벨. Enter로 추가, × 삭제.","s":13,"c":MUTED}]},
    {"space_before":4,"line_spacing":1.05,"runs":[{"t":"현재 34개: "+TAGS[:120]+"…","s":11,"c":GOLD_DK}]},
])
# photos
rect(s, Inches(0.95), Inches(3.65), Inches(11.45), Inches(1.05), fill=WHITE, line=LINE, line_w=1)
textbox(s, Inches(1.2), Inches(3.8), Inches(11.0), Inches(0.85), [
    {"runs":[{"t":"사진  ","s":15,"c":INK,"b":True},{"t":"시술/수술 ‘Photo’ = 카드 썸네일·상세 큰 사진. 홈/로고/메인영상은 코드에 박혀 개발자가 교체.","s":13,"c":MUTED}]},
    {"space_before":4,"runs":[{"t":"환자 전후 사진은 반드시 서면 동의 후에만.","s":12,"c":PINK,"b":True}]},
])
# FAQ
faq = [
    ("저장했는데 안 보여요","‘Visible’ 체크가 꺼졌는지 확인"),
    ("가격 비워도 되나요","네 → ‘Consult’로 표시"),
    ("한글로 적어도 되나요","안 됨 — 영어로(초안은 한글 OK)"),
    ("추천이 안 떠요","‘고민 연결’ 됐는지 확인"),
    ("Why for 설명이 없어요","고민 ‘이유’를 채우면 나옴"),
    ("실수로 삭제","숨김 처리라 복구 가능(개발자 문의)"),
]
rect(s, Inches(0.95), Inches(4.85), Inches(11.45), Inches(1.95), fill=CARD, line=LINE, line_w=1)
lines=[{"runs":[{"t":"자주 묻는 질문","s":15,"c":INK,"b":True}]}]
for q,a in faq:
    lines.append({"space_before":5,"runs":[{"t":"Q. ","s":12,"c":GOLD,"b":True},{"t":q+"  ","s":12,"c":INK,"b":True},{"t":"→ "+a,"s":12,"c":MUTED}]})
textbox(s, Inches(1.2), Inches(5.0), Inches(11.0), Inches(1.7), lines)
nextpage(s)

# ---------------- save ----------------
out_dir = os.path.join(os.path.dirname(__file__), "..", "docs")
os.makedirs(out_dir, exist_ok=True)
target = os.path.join(out_dir, "v3_admin_사용가이드.pptx")
try:
    prs.save(target)
except PermissionError:
    target = target.replace(".pptx", "_new.pptx")
    prs.save(target)
print("생성 완료:", os.path.abspath(target), "·", len(prs.slides.__iter__.__self__._sldIdLst), "슬라이드")
